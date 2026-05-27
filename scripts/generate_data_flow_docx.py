"""Generate Pre-Sales data flow Word doc — dark-theme diagrams + brief explanations."""

from __future__ import annotations

import base64
import io
import shutil
import sys
import zlib
from pathlib import Path

import requests
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

REPO_ROOT = Path(__file__).resolve().parents[1]
OUTPUT = REPO_ROOT / "docs" / "Pre-Sales_Data_Flow_Diagram.docx"
CACHE_DIR = REPO_ROOT / "docs" / "_diagram_cache"

DARK_INIT = "%%{init: {'theme': 'dark'}}%%\n"

# (section title, short explanation, mermaid body — same as first chat message)
SECTIONS: list[tuple[str, str, str]] = [
    (
        "1. End-to-end system data flow",
        "What moves: The browser sends files and JSON over HTTPS (proxied to :8000). The API "
        "writes uploads to disk, runs the orchestrator, and keeps live state in memory while "
        "mirroring to SQLite. LLM calls return structured JSON that becomes requirements, summary, "
        "questions, and eventually the proposal. Email is optional via SMTP or saved to the outbox.",
        """flowchart TB
    subgraph UI["React SPA (:5173)"]
        LOGIN[Login / Register]
        INTAKE[Intake — upload RFP]
        RESULTS[Results / Agents workspace]
        CHAT[Clarification chat]
        SUBMIT[Submit answers + files]
        DRAFT[Draft + Timeline tabs]
    end

    subgraph API["FastAPI (:8000)"]
        AUTH_R["/api/v1/auth/*"]
        RFP_R["/api/v1/rfp/*"]
        JWT[JWT validation]
        ORCH[PipelineOrchestrator]
    end

    subgraph STORE["Persistence"]
        UPLOADS[(data/uploads/)]
        SQLITE[(runs.sqlite3)]
        MEM["RUNS + OUTPUTS caches"]
        CHAT_DB[(chat messages)]
        OUTBOX[(data/email_outbox/)]
    end

    subgraph AI["AI & retrieval"]
        LLM[Groq / Ollama / Azure / mock]
        VS[Vector store memory/chroma]
        AGENTS[Specialist agents]
    end

    subgraph NOTIFY["Notifications"]
        SMTP[SMTP Gmail etc.]
        SCHED[APScheduler reminders]
    end

    LOGIN -->|email + password| AUTH_R
    AUTH_R -->|bcrypt verify| SQLITE
    AUTH_R -->|JWT Bearer| JWT

    INTAKE -->|multipart file + JWT| RFP_R
    JWT --> RFP_R
    RFP_R -->|bytes| UPLOADS
    RFP_R --> ORCH

    ORCH --> AGENTS
    AGENTS -->|chat_json / chat_text| LLM
    ORCH -->|chunks| VS
    ORCH -->|state + OutputBundle| MEM
    MEM <-->|upsert| SQLITE

    RESULTS -->|GET status / output| RFP_R
    RFP_R --> MEM

    CHAT -->|POST chat message| RFP_R
    RFP_R --> AGENTS
    RFP_R --> CHAT_DB

    SUBMIT -->|Form: answers JSON + files| RFP_R
    RFP_R -->|RAG k=8 + clarification context| ORCH
    ORCH -->|final draft| MEM
    RFP_R -->|timeline email| SMTP
    RFP_R -->|fallback copy| OUTBOX

    SCHED -->|3 days before milestone| SMTP
    DRAFT -->|GET output| RFP_R""",
    ),
    (
        "2. Phase 1 — automatic pipeline (upload → 100%, no draft)",
        "Explanation: Phase 1 is fully automatic after upload. Each agent step updates progress, "
        "writes events, and persists to SQLite. Chunks go into the vector store for later RAG. The "
        "pipeline stops at 100% without a proposal so the final draft can use human clarifications.",
        """flowchart LR
    subgraph IN["Input"]
        FILE["RFP file\\nPDF / DOCX / TXT"]
    end

    subgraph P1["PipelineOrchestrator._execute_pipeline"]
        direction TB
        I1["IngestionPlugin\\n20%"]
        GNG{"Go/No-Go\\nempty doc?"}
        I2["UnderstandingAgent\\n40%"]
        TL["build_strict_timeline"]
        I3["SummarizationAgent\\n60%"]
        I4["ScopeGatheringAgent\\n70%"]
        I5["ClarificationAgent\\n80%"]
        DONE["100% COMPLETE\\ndraft = null"]
    end

    subgraph DATA["Data artifacts in PipelineState / OutputBundle"]
        RFP_DOC["RFPDocument\\nraw_text, chunks, sections"]
        REQ["Requirement[]\\nREQ-###"]
        SUM["RFPSummary"]
        SCOPE["PresalesScopeDocument"]
        Q["ClarifyingQuestion[]"]
        TIMELINE["TimelineMilestone[]"]
    end

    subgraph SIDE["Side effects"]
        VS_UP["vector_store.upsert(chunks)"]
        CHAT_SEED["chat_store: opening_message"]
        REM["schedule_project_reminders"]
    end

    FILE --> I1
    I1 --> RFP_DOC
    I1 --> VS_UP
    I1 --> GNG
    GNG -->|abort| ABORT["ABORTED 100%"]
    GNG -->|ok| I2
    I2 --> REQ
    I2 --> TL
    TL --> TIMELINE
    I2 --> I3
    I3 --> SUM
    I3 --> I4
    I4 --> SCOPE
    I4 --> I5
    I5 --> Q
    I5 --> DONE
    DONE --> CHAT_SEED
    DONE --> REM

    DONE -.->|persist| DB[("SQLite + OUTPUTS cache")]""",
    ),
    (
        "3. Phase 2 — clarifications → final draft → email",
        "Explanation: Phase 2 only starts when the user submits clarifications. The system merges "
        "form answers, chat history, and uploaded file text into one prompt context, pulls top-8 "
        "RAG chunks, and generates the integrated proposal. A timeline email fires on success.",
        """sequenceDiagram
    participant U as User
    participant F as Frontend
    participant A as POST /clarifications/resolve
    participant B as clarification_builder
    participant O as Orchestrator
    participant R as RetrievalToolPlugin
    participant V as Vector store
    participant D as ResponseDraftAgent
    participant L as LLM
    participant N as notification_service

    U->>F: Answers + optional PDF/DOCX uploads
    F->>A: Form payload (JSON) + files + JWT

    A->>A: Parse answers, suggested_inputs, notes
    A->>A: Extract text from uploaded files
    A->>B: build_clarification_document(bundle, chat, uploads)
    B-->>A: ClarificationDocument
    A->>B: clarification_to_prompt_context(doc)
    B-->>A: clarification_context string

    A->>O: regenerate_final_draft(run_id, context)
    O->>R: retrieve_context(summary keywords, k=8)
    R->>V: semantic / chunk query
    V-->>R: RAG snippets
    O->>O: merge RAG + clarification_context
    O->>D: draft(..., final_pass=true)
    D->>L: chat_json / prose fallback
    L-->>D: 15-section proposal JSON
    D-->>O: DraftResponse
    O->>O: OUTPUTS + SQLite persist

    A->>N: notify_draft_complete(user, bundle)
    N-->>U: Timeline email (SMTP or outbox)

    A-->>F: final_sections, notification status""",
    ),
    (
        "4. LLM request path (every agent call)",
        "Explanation: All specialist agents share one LLM client. Responses must parse as JSON "
        "into typed models. Groq runs use cooldown delays between steps to avoid TPM limits. On "
        "failure, the run is marked failed with a user-visible error.",
        """flowchart TD
    AGENT["Agent e.g. UnderstandingAgent"]
    PROMPT["System + user prompts\\n+ JSON-only suffix"]
    CLIENT["OpenAICompatibleClient\\nbackend/core/llm.py"]
    BACKEND{"LLM_BACKEND"}

    GROQ["Groq API"]
    OLLAMA["Ollama local"]
    AZURE["Azure OpenAI"]
    MOCK["Mock responses"]

    PARSE["_extract_json\\nfences, repair, truncate"]
    MODEL["Pydantic model\\nRequirementList, RFPSummary, etc."]
    STATE["PipelineState / OutputBundle"]

    AGENT --> PROMPT --> CLIENT --> BACKEND
    BACKEND --> GROQ
    BACKEND --> OLLAMA
    BACKEND --> AZURE
    BACKEND --> MOCK
    GROQ --> PARSE
    OLLAMA --> PARSE
    AZURE --> PARSE
    MOCK --> PARSE
    PARSE -->|ok| MODEL --> STATE
    PARSE -->|fail| ERR["PipelineStatus.FAILED\\nformat_pipeline_error"]""",
    ),
    (
        "5. Persistence & cache hydration",
        "Explanation: Hot path uses in-memory dicts for fast polling (GET /status, WebSocket "
        "stream). SQLite is the source of truth across restarts. Upload binaries stay on disk; "
        "metadata references stored_file UUID names.",
        """flowchart TB
    REQ["API request"]
    RUNS["RUNS dict\\nPipelineState"]
    OUTS["OUTPUTS dict\\nOutputBundle"]
    RS["run_store.py"]
    DB[("data/runs.sqlite3")]
    UP[("data/uploads/{uuid}.ext")]

    REQ --> RUNS
    REQ --> OUTS
    RUNS <-->|upsert_run_state| RS
    OUTS <-->|upsert_run_output| RS
    RS <--> DB

    BOOT["uvicorn lifespan\\nhydrate_caches()"]
    BOOT -->|reload missing| RUNS
    BOOT -->|reload missing| OUTS

    UPLOAD["persist_upload"] --> UP""",
    ),
    (
        "6. Auth data flow",
        "Explanation: Each RFP run is tied to user_id. Users only see their own runs. JWT is "
        "required for every RFP endpoint.",
        """flowchart LR
    REG["POST /auth/register"] --> HASH["bcrypt hash password"]
    HASH --> USERS[("users table")]
    LOG["POST /auth/login"] --> VERIFY["verify password"]
    VERIFY --> USERS
    VERIFY --> JWT_OUT["JWT access token"]
    JWT_OUT --> FE["localStorage / Authorization header"]
    FE --> API["All /api/v1/rfp/* routes"]
    API --> DECODE["get_current_user"]
    DECODE --> ACCESS["assert_run_access\\nuser_id on run"]""",
    ),
]

SUMMARY_ROWS = [
    ("Upload", "POST /rfp/analyze", "File on disk, run_id, async task"),
    ("Phase 1", "Automatic", "Requirements, summary, scope, questions, timeline"),
    ("Human loop", "Chat + form", "Chat messages, clarification document"),
    ("Phase 2", "POST /clarifications/resolve", "Final 15-section draft + timeline email"),
    ("Background", "APScheduler", "Deadline reminder emails"),
]


def _with_dark_theme(mermaid_body: str) -> str:
    return DARK_INIT + mermaid_body.strip()


def _mermaid_ink_url(mermaid_code: str) -> str:
    encoded = base64.urlsafe_b64encode(mermaid_code.encode("utf-8")).decode("ascii")
    return f"https://mermaid.ink/img/{encoded}?type=png&theme=dark&bgColor=!1a1a1a"


def _kroki_deflate_url(mermaid_code: str) -> str:
    compressed = zlib.compress(mermaid_code.encode("utf-8"), 9)
    encoded = base64.urlsafe_b64encode(compressed).decode("ascii")
    return f"https://kroki.io/mermaid/png/{encoded}"


def fetch_diagram_png(mermaid_body: str, index: int, timeout: int = 120) -> bytes:
    mermaid_code = _with_dark_theme(mermaid_body)
    cache_path = CACHE_DIR / f"diagram_{index + 1}_dark.png"
    if cache_path.exists():
        return cache_path.read_bytes()

    sources = (
        ("mermaid.ink", _mermaid_ink_url(mermaid_code)),
        ("kroki", _kroki_deflate_url(mermaid_code)),
    )
    last_error = "no response"
    for name, url in sources:
        try:
            resp = requests.get(url, timeout=timeout, headers={"User-Agent": "presales-docx-gen/2.0"})
            if resp.status_code != 200:
                last_error = f"{name}: HTTP {resp.status_code}"
                continue
            content = resp.content
            if len(content) < 400:
                last_error = f"{name}: response too small"
                continue
            if content[:4] != b"\x89PNG":
                last_error = f"{name}: not PNG"
                continue
            CACHE_DIR.mkdir(parents=True, exist_ok=True)
            cache_path.write_bytes(content)
            return content
        except requests.RequestException as exc:
            last_error = f"{name}: {exc}"

    raise RuntimeError(f"Diagram {index + 1} failed to render ({last_error})")


def _add_explanation(doc: Document, text: str) -> None:
    p = doc.add_paragraph(text)
    p.paragraph_format.space_after = Pt(10)
    for run in p.runs:
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)


def _add_centered_picture(doc: Document, png: bytes, width_inches: float = 6.5) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(io.BytesIO(png), width=Inches(width_inches))


def build_document() -> Document:
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(0.6)
        section.bottom_margin = Inches(0.6)
        section.left_margin = Inches(0.7)
        section.right_margin = Inches(0.7)

    title = doc.add_heading("Pre-Sales AI Agent — Data Flow Diagrams", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for i, (heading, explanation, mermaid_body) in enumerate(SECTIONS):
        doc.add_heading(heading, level=1)
        _add_explanation(doc, explanation)
        png = fetch_diagram_png(mermaid_body, i)
        _add_centered_picture(doc, png)
        if i < len(SECTIONS) - 1:
            doc.add_page_break()

    doc.add_page_break()
    doc.add_heading("Quick summary", level=1)
    _add_explanation(
        doc,
        "High-level stages from upload through automated analysis, human clarification, "
        "final draft generation, and scheduled reminders.",
    )
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text = "Stage"
    hdr[1].text = "Trigger"
    hdr[2].text = "Main data produced"
    for stage, trigger, produced in SUMMARY_ROWS:
        row = table.add_row().cells
        row[0].text = stage
        row[1].text = trigger
        row[2].text = produced

    return doc


def main() -> int:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    if CACHE_DIR.exists():
        shutil.rmtree(CACHE_DIR)

    doc = build_document()
    try:
        doc.save(OUTPUT)
        target = OUTPUT
    except PermissionError:
        target = OUTPUT.with_name(OUTPUT.stem + "_NEW.docx")
        doc.save(target)
        print("Note: close the open .docx in Word, then re-run to overwrite the main file.")

    print(f"Wrote {target} ({len(SECTIONS)} dark-theme diagrams + explanations)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
